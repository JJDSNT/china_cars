from __future__ import annotations

import csv
import re
from collections.abc import Iterable
from datetime import datetime
from pathlib import Path

import duckdb
import pandas as pd
import yaml
from openpyxl import load_workbook

from china_cars.db import database_path
from china_cars.paths import OPS_DIR, RAW_DIR


MONTHS = {
    "Jan": "01",
    "Fev": "02",
    "Mar": "03",
    "Abr": "04",
    "Mai": "05",
    "Jun": "06",
    "Jul": "07",
    "Ago": "08",
    "Set": "09",
    "Out": "10",
    "Nov": "11",
    "Dez": "12",
}
COMEX_NCM_PREFIXES = ("8702", "8703", "8704", "8706", "8707", "8708", "87011")
COMEX_CHINA_CODE = "160"


def _norm_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = " ".join(text.replace("\xa0", " ").split())
    return text.strip()


def _norm_key(value: object) -> str:
    return re.sub(r"[^A-Z0-9]+", " ", _norm_text(value).upper()).strip()


def _load_chinese_brands() -> set[str]:
    path = OPS_DIR / "metadata" / "brand_classification.yml"
    config = yaml.safe_load(path.read_text(encoding="utf-8"))
    return {
        _norm_key(item["brand"])
        for item in config["brands"]
        if item.get("include_in_chinese_cars_scope")
    }


def _iter_anfavea_workbooks() -> Iterable[Path]:
    for year in range(2021, 2027):
        files = sorted((RAW_DIR / "anfavea" / str(year)).glob("*autoveiculos*.xlsx"))
        files += sorted((RAW_DIR / "anfavea" / str(year)).glob("*autove_culos*.xlsx"))
        for path in files:
            name = path.name.lower()
            if "importados_por_empresa" in name or "nacionais_por_empresa" in name:
                continue
            yield path


def extract_anfavea_company_monthly() -> pd.DataFrame:
    records: list[dict[str, object]] = []

    for path in _iter_anfavea_workbooks():
        year_match = re.search(r"(20\d{2})", path.name)
        if not year_match:
            continue
        year = int(year_match.group(1))

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet_name = next((s for s in workbook.sheetnames if "empresa" in s.lower()), None)
        if sheet_name is None:
            continue
        sheet = workbook[sheet_name]

        current_group = None
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            group_candidate = _norm_text(row[1] if len(row) > 1 else None)
            sub_group = _norm_text(row[2] if len(row) > 2 else None)
            company_or_brand = _norm_text(row[3] if len(row) > 3 else None)

            if group_candidate and group_candidate not in {"Total"}:
                current_group = group_candidate

            if not company_or_brand and group_candidate != "Total":
                continue
            if group_candidate == "Total":
                company_or_brand = "Total"
                sub_group = "Total"

            for idx, month_name in enumerate(MONTHS, start=4):
                value = row[idx] if len(row) > idx else None
                if value is None:
                    continue
                try:
                    registrations = int(value)
                except (TypeError, ValueError):
                    continue

                records.append(
                    {
                        "ano": year,
                        "mes": int(MONTHS[month_name]),
                        "ano_mes": f"{year}-{MONTHS[month_name]}",
                        "vehicle_group": current_group,
                        "sub_group": sub_group or None,
                        "company_or_brand": company_or_brand,
                        "company_or_brand_norm": _norm_key(company_or_brand),
                        "emplacamentos": registrations,
                        "source_file": str(path),
                        "source_sheet": sheet_name,
                        "source_row": row_number,
                        "extracted_at": datetime.now().isoformat(timespec="seconds"),
                    }
                )

    return pd.DataFrame.from_records(records)


def build_anfavea_outputs(company_monthly: pd.DataFrame) -> dict[str, pd.DataFrame]:
    chinese_brands = _load_chinese_brands()
    df = company_monthly.copy()

    totals = df[df["company_or_brand_norm"].isin({"TOTAL"})].copy()
    valid_periods = set(totals.loc[totals["emplacamentos"].gt(0), "ano_mes"])
    if valid_periods:
        df = df[df["ano_mes"].isin(valid_periods)].copy()

    df["is_chinese_scope"] = df["company_or_brand_norm"].isin(chinese_brands)

    chinese = df[df["is_chinese_scope"]].copy()
    by_brand = (
        chinese.groupby(
            ["ano", "mes", "ano_mes", "vehicle_group", "company_or_brand"],
            dropna=False,
            as_index=False,
        )["emplacamentos"]
        .sum()
        .sort_values(["ano_mes", "vehicle_group", "company_or_brand"])
    )
    monthly = (
        chinese.groupby(["ano", "mes", "ano_mes"], as_index=False)["emplacamentos"]
        .sum()
        .sort_values("ano_mes")
    )

    total = (
        df[df["company_or_brand_norm"].isin({"TOTAL"})]
        .groupby(["ano", "mes", "ano_mes"], as_index=False)["emplacamentos"]
        .sum()
        .rename(columns={"emplacamentos": "emplacamentos_total_anfavea"})
    )
    monthly = monthly.merge(total, on=["ano", "mes", "ano_mes"], how="left")
    monthly["participacao_total_anfavea"] = (
        monthly["emplacamentos"] / monthly["emplacamentos_total_anfavea"]
    )

    return {
        "raw_anfavea_company_monthly": df,
        "gold_anfavea_chinese_registrations_by_brand_monthly": by_brand,
        "gold_anfavea_chinese_registrations_monthly": monthly,
    }


def _matching_ncm_prefix(code: str) -> str | None:
    for prefix in sorted(COMEX_NCM_PREFIXES, key=len, reverse=True):
        if code.startswith(prefix):
            return prefix
    return None


def _read_ncm_descriptions() -> dict[str, str]:
    path = RAW_DIR / "comex_stat" / "auxiliary" / "NCM.csv"
    if not path.exists():
        return {}

    descriptions: dict[str, str] = {}
    with path.open("r", encoding="latin1", newline="") as file:
        reader = csv.DictReader(file, delimiter=";")
        for row in reader:
            code = row.get("CO_NCM")
            description = row.get("NO_NCM_POR")
            if code and description:
                descriptions[code] = description
    return descriptions


def extract_comex_monthly(flow: str = "IMP") -> pd.DataFrame:
    records: list[pd.DataFrame] = []
    usecols = ["CO_ANO", "CO_MES", "CO_NCM", "CO_UNID", "CO_PAIS", "QT_ESTAT", "KG_LIQUIDO", "VL_FOB"]
    source_dir = RAW_DIR / "comex_stat" / flow.lower()
    descriptions = _read_ncm_descriptions()

    for year in range(2021, 2027):
        path = source_dir / f"{flow}_{year}.csv"
        if not path.exists():
            continue
        for chunk in pd.read_csv(
            path,
            sep=";",
            quotechar='"',
            dtype=str,
            usecols=usecols,
            chunksize=500_000,
        ):
            chunk = chunk[chunk["CO_PAIS"].eq(COMEX_CHINA_CODE)].copy()
            if chunk.empty:
                continue

            chunk["codigo_ncm"] = chunk["CO_NCM"].str.zfill(8)
            chunk["ncm_prefixo_consulta"] = chunk["codigo_ncm"].map(_matching_ncm_prefix)
            chunk = chunk[chunk["ncm_prefixo_consulta"].notna()].copy()
            if chunk.empty:
                continue

            if year == 2026:
                chunk = chunk[chunk["CO_MES"].astype(int).le(6)].copy()

            for col in ("QT_ESTAT", "KG_LIQUIDO", "VL_FOB"):
                chunk[col] = pd.to_numeric(chunk[col], errors="coerce").fillna(0)
            chunk["ano"] = chunk["CO_ANO"].astype(int)
            chunk["mes"] = chunk["CO_MES"].astype(int)
            chunk["ano_mes"] = chunk["ano"].astype(str) + "-" + chunk["CO_MES"].str.zfill(2)
            chunk["fluxo"] = "importacao" if flow == "IMP" else "exportacao"
            chunk["pais"] = "China"
            chunk["descricao_ncm"] = chunk["codigo_ncm"].map(descriptions)
            chunk["source_file"] = str(path)
            records.append(
                chunk[
                    [
                        "ano",
                        "mes",
                        "ano_mes",
                        "fluxo",
                        "pais",
                        "ncm_prefixo_consulta",
                        "codigo_ncm",
                        "descricao_ncm",
                        "CO_UNID",
                        "QT_ESTAT",
                        "KG_LIQUIDO",
                        "VL_FOB",
                        "source_file",
                    ]
                ]
            )

    if not records:
        return pd.DataFrame()

    detail = pd.concat(records, ignore_index=True)
    detail = detail.rename(
        columns={
            "CO_UNID": "codigo_unidade_estatistica",
            "QT_ESTAT": "quantidade_estatistica",
            "KG_LIQUIDO": "kg_liquido",
            "VL_FOB": "valor_fob_usd",
        }
    )
    detail["extracted_at"] = datetime.now().isoformat(timespec="seconds")
    return detail


def build_comex_outputs(detail: pd.DataFrame) -> dict[str, pd.DataFrame]:
    by_ncm = (
        detail.groupby(
            ["ano", "mes", "ano_mes", "fluxo", "pais", "ncm_prefixo_consulta", "codigo_ncm", "descricao_ncm"],
            dropna=False,
            as_index=False,
        )[["quantidade_estatistica", "kg_liquido", "valor_fob_usd"]]
        .sum()
        .sort_values(["ano_mes", "ncm_prefixo_consulta", "codigo_ncm"])
    )
    monthly = (
        detail.groupby(["ano", "mes", "ano_mes", "fluxo", "pais"], as_index=False)[
            ["quantidade_estatistica", "kg_liquido", "valor_fob_usd"]
        ]
        .sum()
        .sort_values("ano_mes")
    )
    by_prefix = (
        detail.groupby(
            ["ano", "mes", "ano_mes", "fluxo", "pais", "ncm_prefixo_consulta"],
            as_index=False,
        )[["quantidade_estatistica", "kg_liquido", "valor_fob_usd"]]
        .sum()
        .sort_values(["ano_mes", "ncm_prefixo_consulta"])
    )
    return {
        "raw_comex_china_automotive_detail": detail,
        "gold_comex_china_automotive_monthly": monthly,
        "gold_comex_china_automotive_by_ncm_monthly": by_ncm,
        "gold_comex_china_automotive_by_prefix_monthly": by_prefix,
    }


def write_duckdb(tables: dict[str, pd.DataFrame]) -> None:
    db_path = database_path()
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with duckdb.connect(str(db_path)) as con:
        for name, df in tables.items():
            con.execute(f"drop table if exists {name}")
            con.register("_df", df)
            con.execute(f"create table {name} as select * from _df")
            con.unregister("_df")


def build_all_tables() -> dict[str, pd.DataFrame]:
    anfavea_company = extract_anfavea_company_monthly()
    tables = build_anfavea_outputs(anfavea_company)

    comex_imp = extract_comex_monthly("IMP")
    tables.update(build_comex_outputs(comex_imp))
    return tables
